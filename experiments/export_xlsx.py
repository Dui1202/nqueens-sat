"""Export a benchmark CSV to an Excel workbook without rerunning solvers."""

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


INTEGER_COLUMNS = {
    "n", "run", "original_variables", "auxiliary_variables",
    "total_variables", "clauses",
}
TIME_COLUMNS = {"build_time", "solve_time", "total_time"}
MEMORY_COLUMNS = {"peak_memory_mb"}


def export_csv_to_xlsx(csv_path: Path) -> Path:
    """Read CSV as the source of truth; preserve blanks, numbers, and literal text."""
    if csv_path.suffix.lower() != ".csv":
        raise ValueError("Expected a .csv input file")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Raw results"
    sheet.freeze_panes = "B2"
    sheet.sheet_view.showGridLines = False

    try:
        with csv_path.open(newline="", encoding="utf-8-sig") as stream:
            reader = csv.reader(stream)
            headers = next(reader, [])
            if not headers or any(not header for header in headers):
                raise ValueError("CSV must contain a nonempty header row")
            sheet.append(headers)
            for row in reader:
                if not row:
                    continue
                values = []
                for header, value in zip(headers, row, strict=True):
                    if value == "":
                        values.append(None)
                    elif header in INTEGER_COLUMNS:
                        values.append(int(value))
                    elif header in TIME_COLUMNS | MEMORY_COLUMNS:
                        values.append(float(value))
                    else:
                        values.append(value)
                sheet.append(values)

        for row in sheet:
            for cell in row:
                cell.font = Font(name="Arial", size=10)
                # Solver messages and labels must remain text, even with '='.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                header = headers[cell.column - 1]
                if cell.row > 1 and header in TIME_COLUMNS:
                    cell.number_format = '0.000000" s"'
                elif cell.row > 1 and header in MEMORY_COLUMNS:
                    cell.number_format = '0.000" MB"'
                elif cell.row > 1 and header in INTEGER_COLUMNS:
                    cell.number_format = "#,##0"
                cell.alignment = Alignment(vertical="top", wrap_text=header == "error")

        for cell in sheet[1]:
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="243746")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 24
        for index, header in enumerate(headers, start=1):
            width = max(len(header) + 3, 16)
            if header == "solver":
                width = 26
            elif header == "error":
                width = 70
            sheet.column_dimensions[get_column_letter(index)].width = width

        if sheet.max_row > 1:
            table = Table(displayName="RawResults", ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showRowStripes=True,
            )
            sheet.add_table(table)
        else:
            sheet.auto_filter.ref = sheet.dimensions

        output_path = csv_path.with_suffix(".xlsx")
        workbook.save(output_path)
        return output_path
    finally:
        workbook.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("csv_path", type=Path, help="Benchmark raw_results.csv file")
    args = parser.parse_args()
    print(export_csv_to_xlsx(args.csv_path))
