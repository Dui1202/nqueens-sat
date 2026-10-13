"""Regression coverage for bounded failures inside benchmark workers."""

import errno
import contextlib
import csv
import io
import json
import multiprocessing as mp
import os
from pathlib import Path
import signal
import subprocess
import sys
import tempfile
import textwrap
import time
import unittest
from unittest.mock import patch

from experiments import runner
from src.baselines.errors import solver_error_status


def kill_worker(*args):
    os.kill(os.getpid(), signal.SIGKILL)


def crash_worker(*args):
    os._exit(17)


def native_oom_exit(*args):
    os.write(2, b"std::bad_alloc\n")
    os._exit(134)


def run_external_child(pid_file, *args):
    child = subprocess.Popen([
        sys.executable, "-c",
        "import json, os, signal, sys, time; from pathlib import Path; "
        "signal.signal(signal.SIGTERM, signal.SIG_IGN); "
        "Path(sys.argv[1]).write_text(json.dumps({'worker': os.getppid(), 'child': os.getpid()})); "
        "time.sleep(30)",
        str(pid_file),
    ])
    child.wait()


class SolverMemoryErrorTests(unittest.TestCase):
    def test_python_and_native_allocation_failures_are_oom(self):
        for error in (
            MemoryError(),
            OSError(errno.ENOMEM, "Cannot allocate memory"),
            RuntimeError("std::bad_alloc"),
            RuntimeError("CPLEX Error 1001: Out of memory."),
            "OpenBLAS error: Memory allocation still failed after 10 retries, giving up.",
            "MemoryError: allocation failed",
        ):
            with self.subTest(error=repr(error)):
                self.assertEqual(solver_error_status(error), "OOM")

    def test_unrelated_errors_and_license_failures_keep_their_status(self):
        for error, expected in (
            (OSError(errno.ENOENT, "No such file or directory"), "ERROR"),
            (RuntimeError("Invalid memory parameter"), "ERROR"),
            (RuntimeError("License expired"), "LICENSE_LIMIT"),
        ):
            with self.subTest(error=str(error)):
                self.assertEqual(solver_error_status(error), expected)


@unittest.skipUnless("fork" in mp.get_all_start_methods(), "Fault injection uses fork")
class WorkerMemoryIsolationTests(unittest.TestCase):
    def assert_next_worker_succeeds(self):
        # This deliberately uses the production spawn context after fault
        # injection has ended, so a fresh solver must actually run.
        result = runner.run_with_timeout(4, "sat", "pairwise", 10)
        self.assertEqual(result["status"], "SAT", result)

    def test_worker_allocation_exceptions_do_not_interrupt_later_runs(self):
        for error in (
            MemoryError("simulated allocation failure"),
            OSError(errno.ENOMEM, "Cannot allocate memory"),
            RuntimeError("std::bad_alloc"),
        ):
            with self.subTest(error=repr(error)):
                with patch.object(runner.mp, "get_context", return_value=mp.get_context("fork")), \
                     patch.object(runner, "benchmark_sat_once", side_effect=error):
                    result = runner.run_with_timeout(4, "sat", "pairwise", 5)
                self.assertEqual(result["status"], "OOM", result)
                self.assertTrue(result["error"])
                self.assert_next_worker_succeeds()

    def test_returned_native_oom_diagnostic_is_preserved(self):
        diagnostic = "CPLEX Error 1001: Out of memory."
        with patch.object(runner.mp, "get_context", return_value=mp.get_context("fork")), \
             patch.object(runner, "benchmark_baseline_once", return_value={
                 "status": "ERROR", "error": diagnostic,
             }):
            result = runner.run_with_timeout(4, "cplex-mip", None, 5)
        self.assertEqual(result["status"], "OOM", result)
        self.assertEqual(result["error"], diagnostic)
        self.assert_next_worker_succeeds()

    def test_unexplained_sigkill_is_not_reported_as_confirmed_oom(self):
        with patch.object(runner.mp, "get_context", return_value=mp.get_context("fork")), \
             patch.object(runner, "benchmark_sat_once", side_effect=kill_worker):
            result = runner.run_with_timeout(4, "sat", "pairwise", 5)
        self.assertEqual(result["status"], "WORKER_KILLED", result)
        self.assertTrue(result["error"])
        self.assert_next_worker_succeeds()

    def test_other_worker_crash_is_an_error_and_next_run_succeeds(self):
        with patch.object(runner.mp, "get_context", return_value=mp.get_context("fork")), \
             patch.object(runner, "benchmark_sat_once", side_effect=crash_worker):
            result = runner.run_with_timeout(4, "sat", "pairwise", 5)
        self.assertEqual(result["status"], "ERROR", result)
        self.assertIn("17", result["error"])
        self.assert_next_worker_succeeds()

    def test_native_oom_exit_preserves_diagnostic_and_next_run_succeeds(self):
        with patch.object(runner.mp, "get_context", return_value=mp.get_context("fork")), \
             patch.object(runner, "benchmark_sat_once", side_effect=native_oom_exit):
            result = runner.run_with_timeout(4, "sat", "pairwise", 5)
        self.assertEqual(result["status"], "OOM", result)
        self.assertIn("std::bad_alloc", result["error"])
        self.assert_next_worker_succeeds()

    @unittest.skipUnless(sys.platform == "linux", "Uses Linux process groups and /proc")
    def test_delayed_group_creation_cannot_escape_timeout_cleanup(self):
        fork_context = mp.get_context("fork")
        allow_group_creation = fork_context.Event()
        original_setsid = os.setsid
        original_killpg = os.killpg
        with tempfile.TemporaryDirectory() as directory:
            pid_file = Path(directory) / "processes.json"

            def delayed_setsid():
                if not allow_group_creation.wait(2):
                    os._exit(18)
                return original_setsid()

            def killpg_with_startup_race(pid, requested_signal):
                try:
                    return original_killpg(pid, requested_signal)
                except ProcessLookupError:
                    if requested_signal == signal.SIGKILL and not allow_group_creation.is_set():
                        # The kernel has already reported a nonexistent group;
                        # now let the worker create it before cleanup resumes.
                        allow_group_creation.set()
                        deadline = time.monotonic() + 1
                        while time.monotonic() < deadline:
                            try:
                                if "child" in json.loads(pid_file.read_text()):
                                    break
                            except (FileNotFoundError, json.JSONDecodeError):
                                pass
                            time.sleep(0.005)
                    raise

            try:
                with patch.object(runner.mp, "get_context", return_value=fork_context), \
                     patch.object(runner.os, "setsid", side_effect=delayed_setsid), \
                     patch.object(runner.os, "killpg", side_effect=killpg_with_startup_race), \
                     patch.object(runner, "benchmark_sat_once", side_effect=run_external_child), \
                     patch.object(runner, "apply_memory_limit"):
                    result = runner.run_with_timeout(
                        pid_file, "sat", "pairwise", 0.02, memory_limit_mb=256,
                    )
                self.assertEqual(result["status"], "TIMEOUT", result)
                processes = json.loads(pid_file.read_text())
                state = Path(f"/proc/{processes['child']}/stat")
                deadline = time.monotonic() + 1
                while state.exists() and state.read_text().split()[2] != "Z" \
                        and time.monotonic() < deadline:
                    time.sleep(0.01)
                self.assertTrue(
                    not state.exists() or state.read_text().split()[2] == "Z",
                    "Child escaped by creating its group after the first SIGKILL attempt",
                )
            finally:
                if pid_file.exists() and pid_file.stat().st_size:
                    for pid in json.loads(pid_file.read_text()).values():
                        try:
                            os.kill(pid, signal.SIGKILL)
                        except ProcessLookupError:
                            pass

    @unittest.skipUnless(sys.platform == "linux", "Uses /proc to distinguish zombies")
    def test_memory_watchdog_stops_external_child_and_next_run_succeeds(self):
        with tempfile.TemporaryDirectory() as directory:
            pid_file = Path(directory) / "processes.json"

            def sampled_rss(process):
                # Block the first RSS sample until the child is ready, before
                # the parent can consume the worker's readiness message. This
                # reproduces an OOM racing process-group setup notification.
                deadline = time.monotonic() + 2
                while time.monotonic() < deadline:
                    try:
                        processes = json.loads(pid_file.read_text())
                        if "worker" in processes and "child" in processes:
                            return 300_000_000
                    except (FileNotFoundError, json.JSONDecodeError):
                        pass
                    time.sleep(0.005)
                self.fail("External child did not become ready within two seconds")

            try:
                with patch.object(runner.mp, "get_context", return_value=mp.get_context("fork")), \
                     patch.object(runner, "benchmark_sat_once", side_effect=run_external_child), \
                     patch.object(runner, "apply_memory_limit"), \
                     patch.object(runner, "process_tree_rss", side_effect=sampled_rss):
                    result = runner.run_with_timeout(
                        pid_file, "sat", "pairwise", 5, memory_limit_mb=256,
                    )
                self.assertEqual(result["status"], "OOM", result)
                self.assertGreaterEqual(result["peak_memory_mb"], 300)
                processes = json.loads(pid_file.read_text())
                for name, pid in processes.items():
                    state = Path(f"/proc/{pid}/stat")
                    deadline = time.monotonic() + 1

                    def still_running():
                        try:
                            return state.read_text().split()[2] != "Z"
                        except (FileNotFoundError, ProcessLookupError):
                            return False

                    while still_running() and time.monotonic() < deadline:
                        time.sleep(0.01)
                    self.assertFalse(
                        still_running(),
                        f"{name} process survived the memory watchdog",
                    )
                self.assert_next_worker_succeeds()
            finally:
                if pid_file.exists() and pid_file.stat().st_size:
                    for pid in json.loads(pid_file.read_text()).values():
                        try:
                            os.kill(pid, signal.SIGKILL)
                        except ProcessLookupError:
                            pass


class WorkerMemoryBudgetTests(unittest.TestCase):
    @unittest.skipUnless(sys.platform == "linux", "Requires Linux address-space limits")
    def test_solver_imports_fit_budget_without_inheriting_large_blas_thread_pool(self):
        script = """
import json, os
from experiments.runner import run_with_timeout
os.environ['OPENBLAS_NUM_THREADS'] = '64'
result = run_with_timeout(8, 'cp-sat', None, 10, memory_limit_mb=512)
print(json.dumps([result['status'], os.environ['OPENBLAS_NUM_THREADS']]))
"""
        completed = subprocess.run(
            [sys.executable, "-c", script], capture_output=True, text=True, timeout=30,
        )
        self.assertEqual(completed.returncode, 0, completed.stderr)
        self.assertEqual(json.loads(completed.stdout.strip().splitlines()[-1]), ["SAT", "64"])

    @unittest.skipUnless(sys.platform == "linux", "Requires Linux address-space limits")
    def test_real_allocation_limit_spares_parent_and_next_run(self):
        # Run in a clean interpreter: other tests import native libraries whose
        # inherited virtual mappings could already exceed this small limit.
        # The attempted allocation is bounded at 320 MiB even if limiting is
        # accidentally removed, and never tries to exhaust the host's memory.
        script = textwrap.dedent("""
            import json
            import multiprocessing as mp
            import resource
            from unittest.mock import patch
            from experiments import runner

            def allocate(*args):
                allocation = bytearray(320 * 1024 * 1024)
                return {"status": "SAT", "total_time": 0}

            if __name__ == "__main__":
                parent_limit = resource.getrlimit(resource.RLIMIT_AS)
                with patch.object(runner.mp, "get_context", return_value=mp.get_context("fork")), \\
                     patch.object(runner, "benchmark_sat_once", side_effect=allocate):
                    failed = runner.run_with_timeout(
                        4, "sat", "pairwise", 10, memory_limit_mb=256,
                    )
                healthy = runner.run_with_timeout(
                    4, "sat", "pairwise", 10, memory_limit_mb=256,
                )
                print(json.dumps({
                    "failed": failed["status"],
                    "healthy": healthy["status"],
                    "parent_unchanged": resource.getrlimit(resource.RLIMIT_AS) == parent_limit,
                }))
        """)
        completed = subprocess.run(
            [sys.executable, "-c", script],
            cwd=Path(__file__).resolve().parents[1],
            capture_output=True, text=True, timeout=30,
        )
        self.assertEqual(completed.returncode, 0, completed.stderr)
        observed = json.loads(completed.stdout.strip().splitlines()[-1])
        self.assertEqual(observed, {
            "failed": "OOM", "healthy": "SAT", "parent_unchanged": True,
        })

    def test_automatic_memory_budget_is_finite_and_bounded(self):
        from experiments.memory import resolve_memory_limit_mb

        limit = resolve_memory_limit_mb()
        self.assertGreater(limit, 0)
        self.assertLessEqual(limit, 4096)
        self.assertLessEqual(resolve_memory_limit_mb(256), 256)


class BenchmarkFailureContinuationTests(unittest.TestCase):
    def test_failed_warmups_and_runs_preserve_the_entire_matrix_in_csv_and_xlsx(self):
        from experiments import benchmark
        from openpyxl import load_workbook

        configurations = [("sat", "pairwise", "test"), ("sat", "sequential", "test")]
        failures = ["OOM", "WORKER_KILLED", "ERROR", "TIMEOUT"]
        outcomes = []
        expected_rows = []
        for index, (n, encoding) in enumerate(
            (n, encoding) for n in (4, 8) for encoding in ("pairwise", "sequential")
        ):
            failure = failures[index]
            # Each failed warm-up must still be followed by both repetitions.
            for status in (failure, failure, "SAT"):
                outcomes.append({
                    **runner.failed_result(status, error=None if status == "SAT" else status),
                    "memory_limit_mb": 256.0,
                })
            expected_rows.extend([(n, encoding, 1, failure), (n, encoding, 2, "SAT")])

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with patch.object(benchmark, "RESULTS_DIR", root), \
                 patch.object(benchmark, "N_VALUES", [4, 8]), \
                 patch.object(benchmark, "WARMUP_RUNS", 1), \
                 patch.object(benchmark, "REPETITIONS", 2), \
                 patch.object(benchmark, "MEMORY_LIMIT_MB", None, create=True), \
                 patch.object(benchmark, "benchmark_configurations", return_value=configurations), \
                 patch.object(benchmark, "run_with_timeout", side_effect=outcomes) as run, \
                 contextlib.redirect_stdout(io.StringIO()):
                benchmark.main()

            csv_files = list(root.glob("*/raw_results.csv"))
            self.assertEqual(len(csv_files), 1)
            with csv_files[0].open(newline="", encoding="utf-8") as stream:
                csv_rows = list(csv.DictReader(stream))
            self.assertEqual(
                [(int(row["n"]), row["encoding"], int(row["run"]), row["status"]) for row in csv_rows],
                expected_rows,
            )
            self.assertTrue(all(float(row["memory_limit_mb"]) == 256 for row in csv_rows))

            workbook = load_workbook(csv_files[0].with_suffix(".xlsx"), read_only=True)
            try:
                values = workbook.active.iter_rows(values_only=True)
                headers = next(values)
                xlsx_rows = [dict(zip(headers, row)) for row in values]
                self.assertEqual(
                    [(row["n"], row["encoding"], row["run"], row["status"]) for row in xlsx_rows],
                    expected_rows,
                )
                self.assertTrue(all(row["memory_limit_mb"] == 256 for row in xlsx_rows))
            finally:
                workbook.close()
            self.assertEqual(run.call_count, 12)
            for call in run.call_args_list:
                self.assertIn("memory_limit_mb", call.kwargs)
                self.assertIsNone(call.kwargs["memory_limit_mb"])


if __name__ == "__main__":
    unittest.main()
